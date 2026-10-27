from pathlib import Path
import json
import shutil
import zipfile

TOPICS = {
 'finance': [
  'Invoice INV-2026-0142. The accounting department requests payment of 4200 AED for office equipment. Payment is due within thirty days. Bank transfer reference is INV-2026-0142.',
  'Quarterly financial budget report. Revenue increased and operating expenses declined. The accounting team forecasts cash flow, payments, invoices and annual profit.',
  'Accounts payable procedure. Review each supplier invoice, validate the purchase order, confirm bank details and record payment in the financial ledger.',
  'Annual financial audit. Reconcile bank statements and verify invoices, expenses and revenue. The accounting team prepares the balance sheet and cash flow report.'
 ],
 'nlp': [
  'Natural language processing research. Sentence embeddings represent text as vectors. Semantic similarity compares documents using cosine distance and transformer representations.',
  'Document clustering experiment. We apply HDBSCAN to sentence transformer embeddings. Related documents share semantic meaning and form groups in vector space.',
  'Text retrieval design. Hybrid search combines keyword matching with semantic embeddings. Query vectors retrieve relevant passages through cosine similarity.',
  'Transformer language model evaluation. Tokenization splits text into subword units. Sentence embeddings support semantic search and unsupervised document clustering.'
 ],
 'cooking': [
  'Tomato pasta recipe. Boil pasta in salted water. Fry garlic in olive oil, add tomatoes and basil, then simmer the sauce. Serve the cooked pasta with cheese.',
  'Vegetable soup recipe. Chop carrots, potatoes and onions. Cook vegetables in stock, add herbs and simmer until tender. Serve hot with bread.',
  'Bread baking recipe. Mix flour, yeast, water and salt. Knead the dough and allow it to rise. Bake in a hot oven until the crust becomes golden.',
  'Roasted vegetable dinner. Toss potatoes and carrots with olive oil, garlic and herbs. Roast in the oven and serve with fresh bread and tomato sauce.'
 ]}

def build(root):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    truth = {'topics': {}, 'duplicates': [], 'queries': []}
    for topic, texts in TOPICS.items():
        for i, text in enumerate(texts):
            name = f'{topic}-{i + 1}.txt'
            (root / name).write_text(text)
            truth['topics'][name] = topic
        truth['queries'].append({'query': {'finance': 'supplier payment and accounting', 'nlp': 'group documents by meaning', 'cooking': 'prepare a meal with vegetables'}[topic],
                                 'relevant': [f'{topic}-{i + 1}.txt' for i in range(4)]})
    shutil.copy2(root / 'finance-1.txt', root / 'invoice-copy.txt')
    truth['topics']['invoice-copy.txt'] = 'finance'
    truth['duplicates'].append(['finance-1.txt', 'invoice-copy.txt', 1])
    # Whitespace changes bytes but not normalized text.
    (root / 'invoice-spaced.txt').write_text(TOPICS['finance'][0].replace(' ', '  '))
    truth['topics']['invoice-spaced.txt'] = 'finance'
    truth['duplicates'] += [['finance-1.txt', 'invoice-spaced.txt', 2], ['invoice-copy.txt', 'invoice-spaced.txt', 2]]
    # Near copy has the same topic and only the due date changes.
    (root / 'invoice-revised.txt').write_text(TOPICS['finance'][0].replace('thirty', 'forty-five'))
    truth['topics']['invoice-revised.txt'] = 'finance'
    truth['duplicates'] += [[name, 'invoice-revised.txt', 3] for name in ['finance-1.txt', 'invoice-copy.txt', 'invoice-spaced.txt']]
    truth['queries'].append({'query': 'INV-2026-0142', 'relevant': ['finance-1.txt', 'invoice-copy.txt', 'invoice-spaced.txt', 'invoice-revised.txt']})
    hard(root, truth)
    return truth


# Harder cases: longer documents, an edited version, a summary, an excerpt and same-topic
# distractors that share vocabulary but are different documents. Queries include paraphrases
# with little word overlap and exact identifiers that only keyword search should find.
AUDIT = ('Annual financial audit report for Falcon Trading LLC, fiscal year 2026. The finance department reviewed all '
         'quarterly accounts and reconciled every bank statement against the general ledger. Total reconciled revenue '
         'reached 1450000 AED against operational expenditures of 820000 AED. Supplier purchase orders and accounts '
         'payable ledgers were verified against delivery notes. Two supplier invoices were paid twice and have been '
         'recovered. Corporate tax provisions and compliance obligations meet regulatory standards. The audit '
         'committee recommends quarterly reconciliation of petty cash and automated three-way matching for invoices. '
         'Reference AUD-7731.')
PAPER = ('This study evaluates sentence embeddings for grouping personal documents without supervision. We encode each '
         'document with a compact transformer, average the passage vectors and cluster them with HDBSCAN. On a '
         'collection of 3200 mixed office files, density clustering recovered 41 topic groups and left 12 percent of '
         'files as noise. Keyword labels from class-based TF-IDF were preferred by annotators over centroid titles. '
         'We also find that hybrid retrieval combining BM25 with dense vectors improves recall for identifier queries '
         'while keeping paraphrase matches. Code name PRJ-5520.')
HARD = {
 'audit-2026.txt': ('finance', AUDIT),
 'audit-2026-final.txt': ('finance', AUDIT.replace('Two supplier invoices were paid twice and have been recovered.',
                                                   'One supplier invoice was paid twice and has been recovered.')
                                          .replace('petty cash', 'petty cash balances')),
 'audit-2026-summary.txt': ('finance', 'Executive summary of the Falcon Trading audit. Revenue of 1450000 AED exceeded '
                            'costs of 820000 AED. Duplicate supplier payments were recovered and tax obligations are met.'),
 'audit-2026-excerpt.txt': ('finance', AUDIT[:AUDIT.index(' Supplier purchase orders')]),
 'audit-harbor-2025.txt': ('finance', 'Annual financial audit report for Harbor Logistics, fiscal year 2025. The auditors '
                           'reviewed quarterly accounts and bank reconciliations. Revenue reached 990000 AED and operating '
                           'costs were 610000 AED. Accounts payable controls were effective and no duplicate payments were '
                           'found. Tax provisions comply with regulations. Reference AUD-4410.'),
 'embeddings-paper.txt': ('nlp', PAPER),
 'embeddings-paper-abstract.txt': ('nlp', 'Abstract. Sentence embeddings and HDBSCAN organise 3200 office files into 41 '
                                   'topic groups without labels, and hybrid BM25 plus dense retrieval helps identifier queries.'),
 'retrieval-notes.txt': ('nlp', 'Notes on neural retrieval. Dense passage retrieval encodes queries and passages with '
                         'separate transformers. Rerankers score the top candidates with cross attention. BM25 remains a '
                         'strong baseline for rare terms and product codes.'),
 'risotto.txt': ('cooking', 'Mushroom risotto. Toast arborio rice in butter, add white wine and ladle in hot stock while '
                 'stirring for eighteen minutes. Fold in parmesan and saute mushrooms, then rest before serving.'),
}


def hard(root, truth):
    for name, (topic, text) in HARD.items():
        (root / name).write_text(text)
        truth['topics'][name] = topic
    truth['duplicates'] += [['audit-2026.txt', 'audit-2026-final.txt', 3],
                            ['audit-2026.txt', 'audit-2026-summary.txt', 4], ['audit-2026-final.txt', 'audit-2026-summary.txt', 4],
                            ['audit-2026.txt', 'audit-2026-excerpt.txt', 4], ['audit-2026-final.txt', 'audit-2026-excerpt.txt', 4],
                            ['embeddings-paper.txt', 'embeddings-paper-abstract.txt', 4]]
    truth['queries'] += [
        {'query': 'AUD-7731', 'relevant': ['audit-2026.txt', 'audit-2026-final.txt']},
        {'query': 'PRJ-5520', 'relevant': ['embeddings-paper.txt']},
        {'query': 'money paid to a vendor two times', 'relevant': ['audit-2026.txt', 'audit-2026-final.txt', 'audit-2026-summary.txt']},
        {'query': 'which files could not be assigned to any group', 'relevant': ['embeddings-paper.txt']},
        {'query': 'slow cooked rice dish with cheese', 'relevant': ['risotto.txt']},
        {'query': 'how are search results ordered for rare product codes', 'relevant': ['retrieval-notes.txt', 'embeddings-paper.txt']},
    ]

def formats(root):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from PIL import Image, ImageDraw, ImageFont
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading('Offline research notes', 0)
    doc.add_paragraph(TOPICS['nlp'][0])
    doc.save(root / 'research.docx')
    wb = Workbook()
    wb.active.append(['Invoice', 'Amount', 'Currency'])
    wb.active.append(['INV-2026-0142', 4200, 'AED'])
    wb.save(root / 'ledger.xlsx')
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide.shapes.title.text = 'Semantic document search'
    slide.placeholders[1].text = TOPICS['nlp'][2]
    ppt.save(root / 'search.pptx')
    (root / 'notes.html').write_text('<html><body><h1>Kitchen notebook</h1><p>' + TOPICS['cooking'][0] + '</p></body></html>')
    (root / 'records.json').write_text(json.dumps({'invoice': 'INV-2026-0142', 'amount': 4200}))
    (root / 'message.eml').write_text('From: sender@example.test\nTo: reader@example.test\nSubject: Offline invoice\nMIME-Version: 1.0\nContent-Type: text/plain; charset=utf-8\n\n' + TOPICS['finance'][0])
    with zipfile.ZipFile(root / 'notes.zip', 'w') as z:
        z.writestr('nested/readme.txt', TOPICS['nlp'][0])
        z.writestr('../escape.txt', 'This archive traversal must never be extracted.')
    image = Image.new('RGB', (1500, 500), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 38) if Path('/System/Library/Fonts/Supplemental/Arial.ttf').exists() else ImageFont.load_default(size=38)
    draw.text((50, 80), 'OFFLINE INVOICE INV-2026-0142\nPayment amount 4200 AED\nPlease pay within thirty days.', fill='black', font=font, spacing=30)
    image.save(root / 'scanned-invoice.png')
    # A text PDF generated with the existing python-docx dependency is not possible;
    # use a tiny standards-compliant PDF builder to avoid adding a runtime extractor.
    content = b'BT /F1 16 Tf 50 750 Td (Offline Invoice INV-2026-0142) Tj 0 -30 Td (Payment amount 4200 AED) Tj ET'
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
       b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
       b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
       b'<< /Length ' + str(len(content)).encode() + b' >>\nstream\n' + content + b'\nendstream']
    data = b'%PDF-1.4\n'
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += str(i).encode() + b' 0 obj\n' + obj + b'\nendobj\n'
    offset = len(data)
    data += b'xref\n0 6\n0000000000 65535 f \n' + b''.join(f'{o:010} 00000 n \n'.encode() for o in offsets[1:])
    data += b'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n' + str(offset).encode() + b'\n%%EOF'
    (root / 'invoice.pdf').write_bytes(data)
    (root / 'wrong-extension.txt').write_bytes(data)
